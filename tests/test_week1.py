"""
tests/test_week1.py -- Week 1 tests (Study A)

Builds a SYNTHETIC Word document with the same kinds of layout as the original document (captions,
an image in a label cell, a merged header, a row of values in parentheses, an appendix
with month labels like 2000.10) and checks that extract_original.py handles them.
Every number in the synthetic document is a placeholder made up by this test
(1.00, 2.00, ...); none comes from the original document.

A fake WindPy module is used to test wind_pull.py without a Wind terminal.

Run from the project root:   python -m unittest discover -s tests -v
"""
import csv
import importlib.util
import io
import shutil
import struct
import sys
import tempfile
import types
import unittest
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "code" / "python" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod            # dataclasses need the module registered
    spec.loader.exec_module(mod)
    return mod


def tiny_png() -> bytes:
    """A valid 1x1 PNG, made in code (stands in for an equation image)."""
    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", zlib.compress(b"\x00\x00")) + chunk(b"IEND", b""))


def add_image(cell, text=""):
    run = cell.paragraphs[0].add_run(text)
    run.add_picture(io.BytesIO(tiny_png()))


def make_synthetic_docx(path: Path, n_months: int = 144) -> None:
    from docx import Document
    doc = Document()
    doc.add_paragraph("CHAPTER 3")
    # Table 3.2-like: variables down the rows, one label is an image
    doc.add_paragraph("Table 3.2 Summarize the first period data")
    t = doc.add_table(rows=1, cols=3)
    for j, h in enumerate(["Variable", "Obs.", "Mean"]):
        t.cell(0, j).text = h
    for k, lab in enumerate(["FER", "GDP", "ER", "AER", "NX", None, "FDI", "FDB"]):
        r = t.add_row().cells
        if lab is None:
            add_image(r[0])
        else:
            r[0].text = lab
        r[1].text = "1"
        r[2].text = f"{k + 1}.00"
    doc.add_paragraph("Note: placeholder note")
    doc.add_paragraph("Table 3.2 showed that the synthetic numbers are placeholders.")
    # Table 4.9-like: merged header, image row label, parenthesis rows, stars
    doc.add_paragraph("Table 4.9 Estimation result of DFER for period 1 and period 2")
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 1).text = "Period 1"
    t.cell(0, 1).merge(t.cell(0, 2))
    t.cell(1, 1).text, t.cell(1, 2).text = "Model I", "Model II"
    rows = [(None, "1.50", ""), ("", "(0.50)", ""), ("DNDF", "", "2.00*"), ("", "", "(1.00)"),
            ("DGDP", "\u22120.0001", "-.0002"), ("", "(0.0003)", "(0.0004)"),
            ("DNX", "3.00", "4.00"), ("", "(1.00)", "(1.00)"),
            ("DFDB", "0.40 **", "0.50\u2020"), ("", "(0.10)", "(0.20 )"),
            ("DFDI", "1,234.50", "5.00"), ("", "(1.00)", "(1.00)"),
            ("_cons", "2.00***", "\u2026")]
    for lab, a, b in rows:
        r = t.add_row().cells
        if lab is None:
            add_image(r[0], "t")
        else:
            r[0].text = lab
        r[1].text, r[2].text = a, b
    # Table 4.11-like text cells
    doc.add_paragraph("Table 4.11 Chow test for structural change")
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 1).text, t.cell(0, 2).text = "F", "P"
    t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = "Model I 1-2", "F (5, 6) = 1.6e+06", "0.0000"
    # Appendix
    doc.add_paragraph("APPENDICES")
    doc.add_paragraph("App. I Monthly data from Jan. 2000 to Dec. 2011")
    header = ["Month", "FER", "GDP", "ER", "AER", "NX", None, "FDI", "FDB", "i"]
    half = n_months // 2
    for part in (range(0, half), range(half, n_months)):
        t = doc.add_table(rows=1, cols=len(header))
        for j, h in enumerate(header):
            if h is None:
                add_image(t.cell(0, j))
            else:
                t.cell(0, j).text = h
        for m in part:
            r = t.add_row().cells
            r[0].text = f"{2000 + m // 12}.{m % 12 + 1}"
            for j in range(1, len(header)):
                r[j].text = f"{j}.00"
            if m == 2:
                r[5].text = "\u22121.00"          # unicode minus
            if m == 3:
                r[2].text = "12,345.00"           # thousands separator
        if part.start == 0:
            r = t.add_row().cells                 # a repeated header row inside the table
            for j, h in enumerate(header):
                r[j].text = h or ""
    doc.add_paragraph("Note: placeholder unit note")
    doc.save(str(path))


class TestParsing(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ex = load("extract_original")

    def test_parse_numeric(self):
        p = self.ex.parse_numeric
        self.assertEqual(p("0.4041 **").value_text, "0.4041")
        self.assertEqual(p("0.4041 **").stars, 2)
        q = p("(0.2196 )")
        self.assertEqual((q.element, q.value_text, q.paren), ("paren", "0.2196", 1))
        self.assertEqual(p("-.0007").value_text, "-.0007")
        self.assertEqual(p("\u22121.5").value_text, "-1.5")
        self.assertEqual(p("1,294.446").value_text, "1294.446")
        self.assertEqual(p("0.6647\u2020").dagger, 1)
        self.assertEqual(p("1.6e+06").element, "value")
        self.assertEqual(p("dU <d<4-dU").element, "text")
        self.assertEqual(p("F (5, 67) =1.0").numbers_in_text, "5;67;1.0")
        self.assertEqual(p("").element, "empty")
        self.assertEqual(p("\u2026").missing, 1)
        self.assertEqual(p(".").element, "text")

    def test_month_labels_stay_text(self):
        f = self.ex.parse_month_label
        self.assertEqual(f("2000.1"), "2000-01")
        self.assertEqual(f("2000.10"), "2000-10")       # would be 2000.1 as a float
        self.assertIsNone(f("2000.13"))
        self.assertIsNone(f("Month"))


class TestExtractionEndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        (self.tmp / "data" / "raw" / "original").mkdir(parents=True)
        self.docx = self.tmp / "data" / "raw" / "original" / "synthetic.docx"
        make_synthetic_docx(self.docx)
        self.ex = load("extract_original")
        self.ex.configure_paths(self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def read(self, rel):
        with open(self.tmp / rel, encoding="utf-8", newline="") as f:
            return list(csv.DictReader(f))

    def test_full_run(self):
        before = self.docx.read_bytes()
        self.assertEqual(self.ex.main([]), 0)
        self.assertEqual(self.docx.read_bytes(), before, "the .docx must not be modified")
        app = self.read("data/interim/original_appendix.csv")
        self.assertEqual(len(app), 144)
        self.assertEqual(app[9]["month_label"], "2000.10")
        self.assertEqual(app[9]["month"], "2000-10")
        self.assertIn("pi", app[0])                      # image header proposed as pi
        self.assertEqual(app[2]["NX"], "-1.00")
        lm = self.read("config/original_label_map.csv")
        self.assertEqual(len(lm), 4)                     # 2 image labels + 2 appendix headers
        props = {(r["table_id"], r["axis"]): r["proposed_label"] for r in lm}
        self.assertEqual(props[("T3.2", "row")], "pi")
        self.assertEqual(props[("T4.9", "row")], "DER_hathat")
        self.assertEqual(props[("APP_I_1", "col")], "pi")
        tg = self.read("data/metadata/original_targets.csv")
        t32 = [t for t in tg if t["table_id"] == "T3.2"]
        self.assertTrue(all(t["table_note"] == "Note: placeholder note" for t in t32))
        paren_rows = [t for t in tg if t["table_id"] == "T4.9" and t["element"] == "paren"]
        self.assertTrue(all(t["row_label"].endswith("(paren)") for t in paren_rows))
        chow = [t for t in tg if t["table_id"] == "T4.11" and t["col_label"] == "F"][0]
        self.assertEqual(chow["numbers_in_text"], "5;6;1.6e+06")
        self.assertEqual(len(self.read("data/metadata/source_ledger.csv")), 1)
        n_tr = len(self.read("data/metadata/transformation_log.csv"))

        # second run: nothing duplicated, map not overwritten
        rows = self.read("config/original_label_map.csv")
        for r in rows:
            r["confirmed_label"], r["confirmed"] = r["proposed_label"], "yes"
        with open(self.tmp / "config" / "original_label_map.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=self.ex.LABEL_MAP_FIELDS, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)
        self.assertEqual(self.ex.main([]), 0)
        self.assertEqual(len(self.read("data/metadata/source_ledger.csv")), 1)
        self.assertEqual(len(self.read("data/metadata/transformation_log.csv")), n_tr)
        self.assertTrue(all(r["confirmed"] == "yes" for r in self.read("config/original_label_map.csv")))
        logs = sorted((self.tmp / "outputs" / "logs").glob("extract_original_*.log"))
        self.assertTrue(logs)
        text = logs[-1].read_text(encoding="utf-8")
        self.assertNotIn("STATUS: FAIL", text)
        self.assertNotIn("still to confirm", text)      # all labels confirmed now

    def test_wrong_row_count_fails(self):
        make_synthetic_docx(self.docx, n_months=140)
        self.assertEqual(self.ex.main([]), 1)
        self.assertFalse((self.tmp / "data" / "metadata" / "source_ledger.csv").exists())


class FakeWindData:
    def __init__(self, times, values, code="X", fields=("CLOSE",)):
        self.ErrorCode, self.Times, self.Data = 0, times, [values]
        self.Codes, self.Fields = [code], list(fields)


class TestWindPull(unittest.TestCase):
    def setUp(self):
        import datetime as dt
        fake = types.ModuleType("WindPy")

        class W:
            def start(self, waitTime=60): return FakeWindData([], [])
            def isconnected(self): return True
            def wsd(self, code, field, begin, end, options=""):
                return FakeWindData([dt.date(2000, 1, 3), dt.date(2000, 1, 4)], [1.0, 2.0], code)
            def edb(self, code, begin, end, options=""):
                return FakeWindData([dt.date(2000, 1, 31)], [3.0], code)
            def wss(self, code, field, options=""):
                d = FakeWindData([], ["placeholder name"], code)
                return d
            def stop(self): pass
        fake.w = W()
        sys.modules["WindPy"] = fake
        self.tmp = Path(tempfile.mkdtemp())
        self.wp = load("wind_pull")

    def tearDown(self):
        sys.modules.pop("WindPy", None)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_pull_and_no_overwrite(self):
        req = self.tmp / "req.csv"
        src = ROOT / "config" / "wind_request.csv"
        with open(src, encoding="utf-8", newline="") as f:
            rows = list(csv.DictReader(f))
            fields = list(rows[0].keys())
        rows[0]["wind_code"], rows[0]["wind_field"], rows[0]["wind_function"] = "TESTCODE", "close", "wsd"
        with open(req, "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
            w.writeheader()
            w.writerows(rows)
        args = ["--request", str(req), "--out-dir", str(self.tmp / "out"), "--log-dir", str(self.tmp / "logs")]
        self.assertEqual(self.wp.main(args), 0)
        files = list((self.tmp / "out").glob("*.xlsx"))
        self.assertEqual(len(files), 1)
        stamp = files[0].stat().st_mtime
        self.assertEqual(self.wp.main(args), 0)            # second run must not overwrite
        self.assertEqual(files[0].stat().st_mtime, stamp)


if __name__ == "__main__":
    unittest.main()
